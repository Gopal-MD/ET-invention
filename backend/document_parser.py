from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path


@dataclass(frozen=True)
class ParsedChunk:
    id: str
    text: str
    source_location: str
    page: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    chunks: list[ParsedChunk]
    parser: str


class DocumentParseError(ValueError):
    pass


def _chunks(text: str, location: str, page: int | None = None) -> list[ParsedChunk]:
    normalized = " ".join(text.split())
    if not normalized:
        return []
    parts = [normalized[index:index + 1200] for index in range(0, len(normalized), 1200)]
    return [
        ParsedChunk(
            id="CHK-" + sha256(f"{location}:{index}:{part}".encode()).hexdigest()[:12].upper(),
            text=part,
            source_location=location,
            page=page,
        )
        for index, part in enumerate(parts)
    ]


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".csv", ".json"}:
            text = content.decode("utf-8-sig")
            chunks = _chunks(text, "text:1")
            parser = "utf8-text"
        elif suffix == ".pdf":
            import fitz
            document = fitz.open(stream=content, filetype="pdf")
            page_chunks: list[ParsedChunk] = []
            page_text: list[str] = []
            for page_number, page in enumerate(document, start=1):
                extracted = page.get_text("text")
                page_text.append(extracted)
                page_chunks.extend(_chunks(extracted, f"page:{page_number}", page_number))
            text = "\n".join(page_text)
            chunks = page_chunks
            parser = "pymupdf"
        elif suffix == ".docx":
            from docx import Document
            document = Document(__import__("io").BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            text = "\n".join(paragraphs)
            chunks = _chunks(text, "paragraphs:1")
            parser = "python-docx"
        else:
            raise DocumentParseError(f"Unsupported file type: {suffix or 'missing extension'}")
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"Could not parse {filename}: {error}") from error
    if not chunks:
        raise DocumentParseError("The document contains no extractable text")
    return ParsedDocument(text=text, chunks=chunks, parser=parser)
